# chatbot/excel_store.py
# ─────────────────────────────────────────────────────────────────────────────
# Save and load student lead data to/from an Excel file (leads.xlsx).
# ─────────────────────────────────────────────────────────────────────────────

import os
import json
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

LEADS_FILE = os.getenv("LEADS_EXCEL_PATH", "leads.xlsx")

COLUMNS = [
    "session_id",
    "timestamp",
    "student_name",
    "phone_number",
    "first_query",
    "chat_summary",
    "course_of_interest",
    "current_education",
    "location",
    "followup_q1",
    "followup_q2",
    "followup_q3",
    "followup_q4",
    "followup_q5",
    "voicebot_status",
    "voicebot_answers",
]


def _get_or_create_workbook() -> tuple[openpyxl.Workbook, object]:
    """Return the workbook and the active sheet, creating the file if needed."""
    if Path(LEADS_FILE).exists():
        wb = openpyxl.load_workbook(LEADS_FILE)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Leads"
        _write_header(ws)
    return wb, ws


def _write_header(ws) -> None:
    """Write a styled header row to a fresh worksheet."""
    header_font   = Font(bold=True, color="FFFFFF", size=11)
    header_fill   = PatternFill(fill_type="solid", fgColor="1E3A5F")
    header_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name.replace("_", " ").title())
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align

    # Set column widths
    widths = {
        1: 36,   # session_id
        2: 22,   # timestamp
        3: 20,   # student_name
        4: 16,   # phone_number
        5: 40,   # first_query
        6: 50,   # chat_summary
        7: 25,   # course_of_interest
        8: 25,   # current_education
        9: 20,   # location
        10: 45,  # followup_q1
        11: 45,  # followup_q2
        12: 45,  # followup_q3
        13: 45,  # followup_q4
        14: 45,  # followup_q5
        15: 18,  # voicebot_status
        16: 50,  # voicebot_answers
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def save_lead(session) -> None:
    """
    Append a lead row to leads.xlsx from a ChatSession object.
    If the file doesn't exist, it will be created with headers.
    """
    wb, ws = _get_or_create_workbook()

    fqs = session.followup_questions or []

    row_data = [
        session.session_id,
        session.started_at,
        session.student_name,
        session.phone_number,
        session.first_query,
        session.chat_summary,
        session.course_of_interest,
        session.current_education,
        session.location,
        fqs[0] if len(fqs) > 0 else "",
        fqs[1] if len(fqs) > 1 else "",
        fqs[2] if len(fqs) > 2 else "",
        fqs[3] if len(fqs) > 3 else "",
        fqs[4] if len(fqs) > 4 else "",
        "pending",   # voicebot_status default
        "",          # voicebot_answers empty initially
    ]

    ws.append(row_data)

    # Alternate row shading
    row_num = ws.max_row
    if row_num % 2 == 0:
        fill = PatternFill(fill_type="solid", fgColor="EBF2FC")
        for col_idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=row_num, column=col_idx).fill = fill

    wb.save(LEADS_FILE)
    print(f"[excel_store] Saved lead for '{session.student_name}' → {LEADS_FILE}")


def load_lead(session_id: str) -> Optional[dict]:
    """
    Load a lead row by session_id. Returns a dict or None if not found.
    Used by Phase 2 voicebot to retrieve lead data before calling.
    """
    if not Path(LEADS_FILE).exists():
        return None

    wb = openpyxl.load_workbook(LEADS_FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == session_id:
            return dict(zip(COLUMNS, row))
    return None


def update_voicebot_status(session_id: str, status: str, answers: dict = None) -> bool:
    """
    Update the voicebot_status and voicebot_answers columns for a given session.
    Used by Phase 2 after a voicebot call completes.
    """
    if not Path(LEADS_FILE).exists():
        return False

    wb = openpyxl.load_workbook(LEADS_FILE)
    ws = wb.active

    status_col  = COLUMNS.index("voicebot_status") + 1
    answers_col = COLUMNS.index("voicebot_answers") + 1

    for row in ws.iter_rows(min_row=2):
        if row[0].value == session_id:
            row[status_col - 1].value  = status
            row[answers_col - 1].value = json.dumps(answers or {})
            wb.save(LEADS_FILE)
            return True
    return False


def list_pending_leads() -> list[dict]:
    """
    Return all leads with voicebot_status='pending'.
    Used by Phase 2 scheduler to find leads to call.
    """
    if not Path(LEADS_FILE).exists():
        return []

    wb = openpyxl.load_workbook(LEADS_FILE)
    ws = wb.active

    pending = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        lead = dict(zip(COLUMNS, row))
        if lead.get("voicebot_status") == "pending" and lead.get("phone_number"):
            pending.append(lead)
    return pending
